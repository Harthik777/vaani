"""Server-side exports: PDF (with Indic shaping when a font is available), Excel RTM, email pack."""
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch, MagicMock
from fastapi.testclient import TestClient
from main import app
from services import config, mailer
import mocks

EN = mocks.MEETING["transcript"]["english"]
HI = mocks.MEETING["transcript"]["hindi"]


def body(lang="english"):
    rules = [{**r, "approved_at": "2026-10-17T10:00:00Z", "approved_by": "Reviewer"} for r in mocks.MEETING["rules"][lang]]
    return {"mode": "business", "rules": rules, "transcript": mocks.MEETING["transcript"][lang], "metadata": {"title": "Spice stock app", "owner": "Indore wholesaler"}}


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.p = patch.object(config, "DATA_DIR", Path(self.tmp.name)); self.p.start()
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close(); self.p.stop(); self.tmp.cleanup()

    def test_pdf_is_a_real_pdf_with_the_requirements(self):
        r = self.client.post("/api/brd/pdf", json=body("english"))
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.headers["content-type"], "application/pdf")
        self.assertIn('filename="Spice-stock-app.pdf"', r.headers["content-disposition"])
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertGreater(len(r.content), 2000)

    def test_pdf_handles_hindi_quotes(self):
        r = self.client.post("/api/brd/pdf", json=body("hindi"))
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.content.startswith(b"%PDF"))

    def test_xlsx_has_one_row_per_requirement_and_extra_sheets(self):
        from openpyxl import load_workbook
        r = self.client.post("/api/brd/xlsx", json={**body("english"), "translations": {"IF stock goes below 100 kg THEN send the reorder alert": "TRANSLATED"}})
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb["Traceability"]
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "ID")
        self.assertEqual(len(rows) - 1, 5)
        self.assertEqual(rows[1][0], "FR-01")
        self.assertIn("Owner", [row[4] for row in rows[1:]])
        self.assertIn("TRANSLATED", [row[3] for row in rows[1:]])
        self.assertIn("Verified against speech", [row[7] for row in rows[1:]])
        self.assertEqual(wb.sheetnames[:3], ["Traceability", "Open questions", "Sections"])

    def test_email_honest_without_smtp_and_sends_pack_when_configured(self):
        r = self.client.post("/api/brd/email", json={**body("english"), "to": ["vendor@example.com"]})
        self.assertEqual(r.status_code, 503)
        self.assertFalse(self.client.get("/api/health").json()["email"])
        sent = {}
        def fake_send(to, subject, text, attachments):
            sent.update(to=to, subject=subject, names=[a[0] for a in attachments], mimes=[a[2] for a in attachments], sizes=[len(a[1]) for a in attachments])
            return {"sent_to": to, "attachments": [a[0] for a in attachments]}
        with patch.object(config, "SMTP_HOST", "smtp.example.com"), patch.object(config, "SMTP_USER", "u"), patch.object(config, "SMTP_PASS", "p"), patch.object(mailer, "send", fake_send):
            self.assertTrue(self.client.get("/api/health").json()["email"])
            r = self.client.post("/api/brd/email", json={**body("english"), "to": ["vendor@example.com", "owner@example.com"], "note": "Please review by Friday."})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(sent["to"], ["vendor@example.com", "owner@example.com"])
        self.assertEqual(sent["subject"], "BRD: Spice stock app")
        self.assertEqual(sent["names"], ["Spice-stock-app.pdf", "Spice-stock-app.docx", "Spice-stock-app-traceability.xlsx"])
        self.assertTrue(all(s > 1000 for s in sent["sizes"]))
        with patch.object(config, "SMTP_HOST", "smtp.example.com"), patch.object(config, "SMTP_USER", "u"), patch.object(config, "SMTP_PASS", "p"):
            self.assertEqual(self.client.post("/api/brd/email", json={**body("english"), "to": ["not-an-email"]}).status_code, 422)

    def test_smtp_send_uses_starttls_and_attaches(self):
        with patch.object(config, "SMTP_HOST", "smtp.example.com"), patch.object(config, "SMTP_PORT", "587"), patch.object(config, "SMTP_USER", "u@example.com"), \
             patch.object(config, "SMTP_PASS", "p"), patch("services.mailer.smtplib.SMTP") as smtp:
            conn = MagicMock(); smtp.return_value.__enter__.return_value = conn
            out = mailer.send(["a@example.com"], "Subject", "Body", [("x.pdf", b"%PDF-1.4", "application/pdf")])
        conn.starttls.assert_called_once(); conn.login.assert_called_once_with("u@example.com", "p")
        msg = conn.send_message.call_args.args[0]
        self.assertEqual(msg["To"], "a@example.com")
        self.assertEqual([p.get_filename() for p in msg.iter_attachments()], ["x.pdf"])
        self.assertEqual(out["attachments"], ["x.pdf"])


if __name__ == "__main__":
    unittest.main()
