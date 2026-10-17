"""Tests for the BRD agent: grounding verifier, interview gaps, conflicts, drafted sections, exports and Sarvam adapters."""
import asyncio
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch
import httpx
from fastapi.testclient import TestClient
from main import app
from services import config, llm, sarvam, verify
import mocks

EN = mocks.MEETING["transcript"]["english"]
HI = mocks.MEETING["transcript"]["hindi"]


def live_llm():
    return [patch.object(config, "USE_REAL_LLM", True), patch.object(config, "LLM_API_KEY", "test"), patch.object(config, "AGENT_LLM", "main")]


class Patched(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.patches = [patch.object(config, "DATA_DIR", Path(self.tmp.name)), patch.object(config, "USE_REAL_SARVAM", False), patch.object(config, "USE_REAL_LLM", False)]
        for p in self.patches:
            p.start()
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        for p in self.patches:
            p.stop()
        self.tmp.cleanup()

    def with_patches(self, *extra):
        for p in extra:
            p.start()
            self.addCleanup(p.stop)


class VerifierTests(unittest.TestCase):
    def test_numbers_are_canonicalised_across_scripts_and_grouping(self):
        self.assertEqual(verify.numbers("₹2,400 per kg"), ["2400"])
        self.assertEqual(verify.numbers("नमी ९.५% से ज्यादा"), ["9.5"])
        self.assertEqual(verify.numbers("20,000 rupees and 1,50,000"), ["20000", "150000"])
        self.assertEqual(verify.numbers("sizes 3,4"), ["3", "4"])
        self.assertEqual(verify.numbers("10.50 and 07"), ["10.5", "7"])
        self.assertEqual(verify.numbers("0.5% and 0 and 2.0"), ["0.5", "0", "2"])

    def test_english_number_words_cannot_bypass_the_check(self):
        self.assertEqual(verify.numbers("below fifty kg"), ["50"])
        self.assertEqual(verify.numbers("twenty thousand rupees, one hundred and five units"), ["20000", "105"])
        self.assertEqual(verify.numbers("one of the staff, a lakh of thanks, often"), [])
        self.assertEqual(verify.numbers("1.5 lakh"), ["1.5"])
        rule = {"title": "Alert", "condition": "stock below fifty kg", "action": "alert", "source_quote": "स्टॉक पचास किलो से कम"}
        self.assertEqual(verify.verify_rule(rule, "जब स्टॉक पचास किलो से कम हो")["status"], "unsupported_number")
        rule["numbers"] = [{"value": "50", "spoken": "पचास किलो"}]
        self.assertEqual(verify.verify_rule(rule, "जब स्टॉक पचास किलो से कम हो")["status"], "needs_confirmation")
        english = {"title": "Hold", "condition": "x", "action": "pause for 3 days", "source_quote": "hold purchases for three days"}
        self.assertEqual(verify.verify_rule(english, "I hold purchases for three days")["status"], "verified")

    def test_number_statuses(self):
        rule = {"title": "Hold", "condition": "dollar spikes", "action": "pause purchases for 72 hours", "source_quote": "I hold purchases for 3 days"}
        self.assertEqual(verify.verify_rule(rule, "x I hold purchases for 3 days y")["status"], "unsupported_number")
        rule = {"title": "Hold", "condition": "dollar spikes", "action": "pause for 3 days", "source_quote": "तीन दिन खरीद रोकता हूँ",
                "numbers": [{"value": "3", "spoken": "तीन दिन"}]}
        check = verify.verify_rule(rule, "मैं तीन दिन खरीद रोकता हूँ")
        self.assertEqual(check["status"], "needs_confirmation")
        self.assertEqual(check["numbers"][0], {"value": "3", "status": "interpreted", "spoken": "तीन दिन"})
        # A claimed spoken form that is not in the quote does not count.
        rule["numbers"] = [{"value": "3", "spoken": "invented words"}]
        self.assertEqual(verify.verify_rule(rule, "मैं तीन दिन खरीद रोकता हूँ")["status"], "unsupported_number")
        self.assertEqual(verify.verify_rule({**rule, "source_quote": "not said"}, "मैं तीन दिन")["status"], "rejected")

    def test_near_exact_quotes_are_repaired_to_real_text_only(self):
        line = "दोबारा ऑर्डर का अलर्ट तब आना चाहिए जब स्टॉक सौ किलो से कम हो"
        tx = "[Owner] something else\n[Manager] " + line + ", क्योंकि त्योहार में देर होती है।"
        self.assertEqual(verify.repair_quote(line.replace("दोबारा", "दबारा"), tx), line)
        self.assertIsNone(verify.repair_quote(line.replace("सौ", "पचास"), tx))  # a changed number word is not a typo
        self.assertIsNone(verify.repair_quote("completely different words that were never said", tx))

    def test_items_suggestions_cannot_carry_numbers(self):
        self.assertEqual(verify.verify_item({"text": "Staff may resist change", "source_quote": None}, EN)[0]["origin"], "suggested")
        self.assertIn("nobody said", verify.verify_item({"text": "Budget of 5 lakh", "source_quote": None}, EN)[1])
        self.assertIn("not found", verify.verify_item({"text": "Track stock", "source_quote": "invented"}, EN)[1])
        kept, _ = verify.verify_item({"text": "Alert below 50 kg", "source_quote": "When cardamom stock falls below 50 kg"}, EN)
        self.assertEqual(kept["numbers"][0]["status"], "exact")


class MeetingDemoTests(Patched):
    def test_meeting_rules_are_grounded_in_both_languages(self):
        for lang, text in (("english", EN), ("hindi", HI)):
            with self.subTest(lang=lang):
                body = self.client.post("/api/extract", json={"transcript": text, "mode": "business", "demo": True}).json()
                self.assertEqual(len(body["rules"]), 5)
                for r in body["rules"]:
                    self.assertIn(r["source_quote"], text)
                    self.assertEqual(r["verification"]["status"], "verified", r["id"])
                self.assertEqual({r["speaker"] for r in body["rules"]}, {"Owner", "Manager"})

    def test_demo_extraction_survives_appended_interview_answers_but_not_edits(self):
        appended = EN + "\n[Owner] About 12 people will use it."
        self.assertEqual(self.client.post("/api/extract", json={"transcript": appended, "demo": True}).status_code, 200)
        self.assertEqual(self.client.post("/api/extract", json={"transcript": EN.replace("50", "60"), "demo": True}).status_code, 422)

    def test_demo_gaps_ask_only_what_is_missing(self):
        body = self.client.post("/api/agent/gaps", json={"transcript": EN, "lang": "english", "demo": True}).json()
        self.assertEqual(body["completeness"], 38)
        self.assertEqual([q["key"] for q in body["questions"]], ["users", "timeline", "budget", "success", "approver"])
        covered = [i for i in body["items"] if i["status"] == "covered"]
        self.assertTrue(all(i["source_quote"] in EN for i in covered))
        self.assertIn("users", body["sample_answers"])
        answered = self.client.post("/api/agent/gaps", json={"transcript": EN, "demo": True, "answered": ["users", "timeline"]}).json()
        self.assertEqual(answered["completeness"], 63)
        self.assertNotIn("users", [q["key"] for q in answered["questions"]])

    def test_demo_conflict_between_owner_and_manager(self):
        rules = mocks.MEETING["rules"]["english"]
        body = self.client.post("/api/agent/conflicts", json={"rules": rules, "transcript": EN, "demo": True}).json()
        self.assertEqual(body["conflicts"][0]["rule_ids"], ["r1", "r4"])
        single = self.client.post("/api/agent/conflicts", json={"rules": rules[:2], "transcript": EN, "demo": True}).json()
        self.assertEqual(single["conflicts"], [])


class LiveAgentTests(Patched):
    def test_coverage_claims_need_real_quotes(self):
        answer = {"items": [
            {"key": "objective", "covered": True, "source_quote": "track stock and orders"},
            {"key": "budget", "covered": True, "source_quote": "budget is 5 lakh"},  # invented
            {"key": "timeline", "covered": False, "question": "When do you need the stock app?"}]}
        self.with_patches(*live_llm(), patch.object(llm, "complete_json", new=AsyncMock(return_value=answer)))
        body = self.client.post("/api/agent/gaps", json={"transcript": EN, "lang": "english"}).json()
        status = {i["key"]: i["status"] for i in body["items"]}
        self.assertEqual(status["objective"], "covered")
        self.assertEqual(status["budget"], "missing")
        self.assertIn("When do you need the stock app?", [q["question"] for q in body["questions"]])

    def test_gap_failure_falls_back_to_full_checklist(self):
        self.with_patches(*live_llm(), patch.object(llm, "complete_json", new=AsyncMock(side_effect=RuntimeError("down"))))
        body = self.client.post("/api/agent/gaps", json={"transcript": EN, "answered": ["users"]}).json()
        self.assertEqual(len(body["questions"]), 7)
        self.assertIn("warning", body)

    def test_conflicts_from_llm_are_validated(self):
        rules = mocks.MEETING["rules"]["english"]
        answer = {"conflicts": [{"rule_ids": ["r1", "r4"], "topic": "threshold", "explanation": "50 vs 100"},
                                {"rule_ids": ["r1", "r2"], "topic": "same speaker", "explanation": "x"},
                                {"rule_ids": ["r1", "r99"], "topic": "unknown", "explanation": "x"},
                                {"rule_ids": ["r4", "r1"], "topic": "duplicate", "explanation": "x"}]}
        self.with_patches(*live_llm(), patch.object(llm, "complete_json", new=AsyncMock(return_value=answer)))
        body = self.client.post("/api/agent/conflicts", json={"rules": rules, "transcript": EN}).json()
        self.assertEqual([c["rule_ids"] for c in body["conflicts"]], [["r1", "r4"]])

    def test_drafted_sections_keep_only_grounded_content(self):
        rules = [{**r, "approved_at": "t", "approved_by": "Owner"} for r in mocks.MEETING["rules"]["english"]]
        answer = {"sections": {
            "objectives": [{"text": "Stop losing orders kept in notebooks", "source_quote": "everything is written in notebooks and we lose orders"},
                           {"text": "Cut losses by 40%", "source_quote": "we lose orders"}],
            "risks": [{"text": "Staff may need training on the new app", "source_quote": None},
                      {"text": "Budget may exceed 2 lakh", "source_quote": None}],
            "constraints": [{"text": "Must run on a server we never mentioned", "source_quote": "we have a server"}],
            "made_up_section": [{"text": "x", "source_quote": None}]},
            "acceptance": [{"rule_id": "r1", "text": "Given stock is 49 kg when saved then an alert is shown"},
                           {"rule_id": "r1", "text": "Given stock falls below 50 kg then the owner is alerted"}]}
        self.with_patches(*live_llm(), patch.object(llm, "complete_json", new=AsyncMock(return_value=answer)))
        body = self.client.post("/api/brd/draft", json={"rules": rules, "transcript": EN, "mode": "business"}).json()
        self.assertEqual([i["text"] for i in body["sections"]["objectives"]], ["Stop losing orders kept in notebooks"])
        self.assertEqual([i["origin"] for i in body["sections"]["risks"]], ["suggested"])
        self.assertEqual(body["sections"]["constraints"], [])
        self.assertEqual(len(body["acceptance"]), 1)
        reasons = " | ".join(r["reason"] for r in body["rejected"])
        for fragment in ("40", "nobody said", "not found", "49"):
            self.assertIn(fragment, reasons)

    def test_draft_requires_rules_and_llm(self):
        self.assertEqual(self.client.post("/api/brd/draft", json={"rules": [], "transcript": EN}).status_code, 422)
        rule = {**mocks.MEETING["rules"]["english"][0], "approved_at": "t", "approved_by": "o"}
        with patch.object(config, "USE_REAL_LLM", False):
            self.assertEqual(self.client.post("/api/brd/draft", json={"rules": [rule], "transcript": EN}).status_code, 503)


class DocumentTests(Patched):
    def body(self):
        rules = [{**r, "approved_at": "2026-10-16T10:00:00Z", "approved_by": "Owner"} for r in mocks.MEETING["rules"]["english"]]
        transcript = EN + "\n[Owner] About 12 people will use it: me, the manager and 10 staff members."
        return {"mode": "business", "rules": rules, "transcript": transcript, "metadata": {"title": "Stock app"},
                "answers": [{"key": "users", "answer": "About 12 people will use it: me, the manager and 10 staff members."}],
                "coverage": [{"key": "objective", "covered": True}], "conflicts": mocks.MEETING["conflicts"],
                "sections": {"risks": [{"text": "Invented 9 risks", "source_quote": None}]}}

    def test_full_brd_structure(self):
        doc = self.client.post("/api/brd", json=self.body()).json()
        self.assertEqual(doc["business_rules"][0]["req_id"], "FR-01")
        self.assertEqual(doc["business_rules"][0]["verification"]["status"], "verified")
        self.assertIn("12 people", doc["sections"]["stakeholders"][0]["text"])
        self.assertEqual(doc["sections"]["risks"], [])
        self.assertTrue(doc["rejected"])
        self.assertTrue(any(q.startswith("Budget") for q in doc["open_questions"]))
        self.assertTrue(any("Conflict to resolve" in q for q in doc["open_questions"]))
        self.assertFalse(any(q.startswith("Users") for q in doc["open_questions"]))
        self.assertNotIn("FastAPI", str(doc))

    def test_docx_export(self):
        response = self.client.post("/api/brd/docx", json=self.body())
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"PK"))
        self.assertIn("Stock-app.docx", response.headers["content-disposition"])
        import io
        import docx
        text = "\n".join(p.text for p in docx.Document(io.BytesIO(response.content)).paragraphs)
        self.assertIn("Open questions", text)
        self.assertIn("Verification report", text)

    def test_app_generation_falls_back_honestly(self):
        body = self.body()
        with patch.object(config, "USE_REAL_LLM", False):
            self.assertEqual(self.client.post("/api/app/generate", json=body).json()["generator"], "template")
        self.with_patches(*live_llm())
        with patch.object(llm, "complete_text", new=AsyncMock(return_value="Sure! here is your app")):
            result = self.client.post("/api/app/generate", json=body).json()
            self.assertEqual(result["generator"], "template")
            self.assertIn("failed", result["note"])
        html = "<!doctype html><html><head></head><body>FR-01</body></html>"
        with patch.object(llm, "complete_text", new=AsyncMock(return_value=html)) as call:
            result = self.client.post("/api/app/generate", json=body).json()
            self.assertEqual(result["html"], html)
            self.assertTrue(result["generator"].startswith("ai:"))
            self.assertIn("FR-01", call.call_args.args[0])


class RequirementsAppTests(Patched):
    def test_simulator_is_built_from_the_brd_and_cannot_be_script_injected(self):
        from services import microapps
        html = microapps.requirements_app("Stock </script><img src=x onerror=alert(1)>", mocks.MEETING["rules"]["english"], ["Budget?"])
        self.assertEqual(html.count("</script>"), 1)
        self.assertIn("stock goes below 100 kg", html)
        self.assertNotIn("Scrap", html)

    def test_business_fallback_app_uses_requirements_not_a_canned_template(self):
        rules = [{**r, "approved_at": "t", "approved_by": "o"} for r in mocks.MEETING["rules"]["english"]]
        with patch.object(config, "USE_REAL_LLM", False):
            body = self.client.post("/api/app/generate", json={"mode": "business", "rules": rules, "transcript": EN}).json()
        self.assertEqual(body["generator"], "template")
        self.assertIn("Low-stock reorder alert", body["html"])
        self.assertNotIn("Forex", body["html"])


class AppLanguageTests(Patched):
    def body(self, language):
        rules = [{**r, "approved_at": "t", "approved_by": "o"} for r in mocks.MEETING["rules"]["english"]]
        return {"mode": "business", "rules": rules, "transcript": EN, "app_language": language}

    def test_simulator_is_translated_with_sarvam(self):
        async def fake_translate(text, target, source="english", client=None):
            return "HI:" + text
        self.with_patches(patch.object(config, "USE_REAL_LLM", False), patch.object(config, "USE_REAL_SARVAM", True),
                          patch.object(config, "SARVAM_API_KEY", "k"), patch.object(sarvam, "translate", new=fake_translate))
        body = self.client.post("/api/app/generate", json=self.body("hindi")).json()
        self.assertEqual(body["language"], "hindi")
        self.assertIn('"lang": "hi"', body["html"])
        self.assertIn("HI:Try a value", body["html"])
        self.assertIn("HI:Low-stock reorder alert", body["html"])
        self.assertIn('"condition": "stock goes below 100 kg"', body["html"])  # thresholds still parse the English rule

    def test_without_sarvam_the_app_says_it_stays_english(self):
        self.with_patches(patch.object(config, "USE_REAL_LLM", False))
        body = self.client.post("/api/app/generate", json=self.body("tamil")).json()
        self.assertEqual(body["language"], "english")
        self.assertIn("shown in English", body["note"])

    def test_ai_app_prompt_asks_for_the_chosen_language(self):
        html = "<!doctype html><html><head></head><body>FR-01</body></html>"
        self.with_patches(*live_llm())
        with patch.object(llm, "complete_text", new=AsyncMock(return_value=html)) as call:
            body = self.client.post("/api/app/generate", json=self.body("tamil")).json()
        self.assertEqual(body["language"], "tamil")
        self.assertIn("in Tamil, in its native script", call.call_args.args[0])

    def test_word_export_is_bilingual_when_translations_are_given(self):
        import io
        import docx
        rules = [{**r, "approved_at": "t", "approved_by": "o"} for r in mocks.MEETING["rules"]["english"]]
        rule_text = "IF stock goes below 100 kg THEN send the reorder alert"
        response = self.client.post("/api/brd/docx", json={"mode": "business", "rules": rules, "transcript": EN, "translations": {rule_text: "TRANSLATED-RULE"}})
        document = docx.Document(io.BytesIO(response.content))
        cells = "\n".join(c.text for t in document.tables for row in t.rows for c in row.cells)
        self.assertIn("TRANSLATED-RULE", cells)


class SarvamAdapterTests(Patched):
    def fake(self, payload):
        client = AsyncMock()
        client.__aenter__.return_value = client
        client.post.return_value = httpx.Response(200, json=payload, request=httpx.Request("POST", "https://api.sarvam.ai/x"))
        return client

    def test_split_text_respects_limits(self):
        text = ("This is a sentence. " * 400).strip()
        pieces = sarvam.split_text(text, 2500)
        self.assertTrue(all(len(p) <= 2500 for p in pieces))
        self.assertEqual(" ".join(pieces).split(), text.split())
        self.assertEqual(sarvam.split_text("x" * 30, 10), ["x" * 10] * 3)

    def test_tts_uses_documented_fields(self):
        fake = self.fake({"audios": ["UklGRg=="]})
        self.with_patches(patch.object(config, "USE_REAL_SARVAM", True), patch.object(config, "SARVAM_API_KEY", "k"), patch("services.sarvam.httpx.AsyncClient", return_value=fake))
        body = self.client.post("/api/tts", json={"text": "नमस्ते", "lang": "hindi"}).json()
        self.assertEqual(body["audios"], ["UklGRg=="])
        sent = fake.post.call_args.kwargs["json"]
        self.assertEqual(sent["language_code"], "hi-IN")
        self.assertNotIn("target_language_code", sent)
        self.assertEqual(sent["model"], config.SARVAM_TTS_MODEL)
        self.assertEqual(fake.post.call_args.kwargs["headers"], {"api-subscription-key": "k"})

    def test_translate_uses_documented_fields(self):
        sarvam._translations.clear()
        fake = self.fake({"translated_text": "आप कब तक चाहते हैं?"})
        self.with_patches(patch.object(config, "USE_REAL_SARVAM", True), patch.object(config, "SARVAM_API_KEY", "k"), patch("services.sarvam.httpx.AsyncClient", return_value=fake))
        body = self.client.post("/api/translate", json={"texts": ["When do you need it?"], "target": "hindi"}).json()
        self.assertEqual(body["translations"], ["आप कब तक चाहते हैं?"])
        sent = fake.post.call_args.kwargs["json"]
        self.assertEqual((sent["input"], sent["source_language_code"], sent["target_language_code"], sent["model"]),
                         ("When do you need it?", "en-IN", "hi-IN", "sarvam-translate:v1"))

    def test_sarvam_chat_uses_json_mode(self):
        fake = self.fake({"choices": [{"message": {"content": '{"ok": true}'}}]})
        self.with_patches(patch.object(config, "SARVAM_API_KEY", "k"), patch("services.sarvam.httpx.AsyncClient", return_value=fake))
        self.assertEqual(asyncio.run(sarvam.chat_json("prompt")), '{"ok": true}')
        sent = fake.post.call_args.kwargs["json"]
        self.assertEqual(sent["response_format"], {"type": "json_object"})
        self.assertEqual(sent["model"], config.SARVAM_LLM_MODEL)
        self.assertTrue(fake.post.call_args.args[0].endswith("/v1/chat/completions"))

    def test_voice_routes_are_honest_when_unconfigured(self):
        self.assertEqual(self.client.post("/api/tts", json={"text": "hi", "lang": "hindi"}).status_code, 503)
        self.assertEqual(self.client.post("/api/translate", json={"texts": ["hi"], "target": "hindi"}).status_code, 503)


if __name__ == "__main__":
    unittest.main()


class LiveDemoTests(Patched):
    """With a model configured, the meeting demo runs through the real agent and extraction (what the audience sees is live);
    the fixture is the safety net when a call fails, and the only path when no model is configured."""
    KEY, QUOTE = next(iter(mocks.MEETING["coverage"]["english"].items()))
    COVER = {"items": [{"key": KEY, "covered": True, "source_quote": QUOTE}, {"key": "users", "covered": False, "question": "Who will use it?"}]}

    def test_demo_gap_check_uses_the_live_agent_and_falls_back_to_the_fixture(self):
        self.with_patches(*live_llm())
        with patch.object(llm, "agent_json", new=AsyncMock(return_value=(self.COVER, "gemini:fake"))):
            body = self.client.post("/api/agent/gaps", json={"transcript": EN, "lang": "english", "demo": True}).json()
        self.assertEqual(body["brain"], "gemini:fake")
        self.assertIn("users", body["sample_answers"])  # the demo's sample answers stay available
        status = {i["key"]: i["status"] for i in body["items"]}
        self.assertEqual((status[self.KEY], status["users"]), ("covered", "missing"))
        with patch.object(llm, "agent_json", new=AsyncMock(side_effect=RuntimeError("429"))):
            body = self.client.post("/api/agent/gaps", json={"transcript": EN, "lang": "english", "demo": True}).json()
        self.assertEqual(body["brain"], "demo fixture")
        self.assertEqual(body["completeness"], 38)
        with patch.object(config, "DEMO_LIVE", False), patch.object(llm, "agent_json", new=AsyncMock(return_value=(self.COVER, "gemini:fake"))) as fake:
            body = self.client.post("/api/agent/gaps", json={"transcript": EN, "lang": "english", "demo": True}).json()
        self.assertEqual(body["brain"], "demo fixture")
        fake.assert_not_called()

    def test_demo_extraction_uses_the_live_model_and_falls_back_to_the_fixture(self):
        self.with_patches(*live_llm())
        answer = {"rules": [{"id": "r1", "title": "Owner threshold", "condition": "cardamom stock falls below 50 kg", "action": "alert the owner",
                             "source_quote": "When cardamom stock falls below 50 kg, the app should alert me to reorder."},
                            {"id": "r2", "title": "Manager threshold", "condition": "stock goes below 100 kg", "action": "send the reorder alert",
                             "source_quote": "The reorder alert should come when stock goes below 100 kg"}],
                  "conflicts": [{"rule_ids": ["r1", "r2"], "topic": "threshold", "explanation": "50 vs 100"}], "review_notes": []}
        with patch.object(llm, "complete_json", new=AsyncMock(return_value=answer)):
            body = self.client.post("/api/extract", json={"transcript": EN, "mode": "business", "demo": True}).json()
        self.assertTrue(body["demo"])
        self.assertNotEqual(body["brain"], "demo fixture")
        self.assertEqual([r["speaker"] for r in body["rules"]], ["Owner", "Manager"])
        self.assertEqual(body["conflicts"][0]["rule_ids"], ["r1", "r2"])
        with patch.object(llm, "complete_json", new=AsyncMock(side_effect=RuntimeError("503"))):
            body = self.client.post("/api/extract", json={"transcript": EN, "mode": "business", "demo": True}).json()
        self.assertEqual(body["brain"], "demo fixture")
        self.assertEqual(len(body["rules"]), 5)
        # Other demo transcripts (not the meeting) and edited meeting text stay on the fixture rules.
        with patch.object(llm, "complete_json", new=AsyncMock(return_value=answer)) as fake:
            self.assertEqual(self.client.post("/api/extract", json={"transcript": EN.replace("50", "60"), "demo": True}).status_code, 422)
        fake.assert_not_called()
