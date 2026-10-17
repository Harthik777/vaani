"""BRD -> .docx (Word) for the IT vendor. Takes the compiled document from compiler.document()."""
from io import BytesIO
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ORIGIN_LABEL = {"spoken": "Spoken", "suggested": "AI suggestion: confirm"}


def _fonts(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:cs"), "Nirmala UI")  # Devanagari, Tamil, Telugu, Bengali glyphs
    fonts.set(qn("w:eastAsia"), "Nirmala UI")


def _note(paragraph, text):
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x5F, 0x6D)


def _translation(doc, text, translations):
    """Bilingual export: the Sarvam translation sits under the English line, which stays the reference."""
    translated = (translations or {}).get(text)
    if translated:
        run = doc.add_paragraph().add_run(translated)
        run.font.color.rgb = RGBColor(0x4B, 0x3F, 0xA8)


def build_docx(brd: dict, translations: dict | None = None) -> bytes:
    doc = Document()
    _fonts(doc)
    doc.add_heading(brd.get("project_name", "Business Requirements Document"), 0)
    p = doc.add_paragraph(brd.get("document_type", ""))
    _note(p, "  ·  " + brd.get("status", ""))
    if brd.get("demo"):
        doc.add_paragraph("DEMONSTRATION: sample content.")
    meta = {k: v for k, v in (brd.get("metadata") or {}).items() if v and k not in {"ticketSummary", "review_notes"}}
    if meta:
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light List Accent 1"
        for k, v in meta.items():
            row = table.add_row().cells
            row[0].text, row[1].text = k.replace("_", " ").title(), str(v)

    titles = brd.get("section_titles", {})
    number = 0
    for key, items in (brd.get("sections") or {}).items():
        number += 1
        doc.add_heading(f"{number}. {titles.get(key, key)}", 1)
        if not items:
            _note(doc.add_paragraph(), "Not captured yet.")
        for item in items:
            para = doc.add_paragraph(item["text"], style="List Bullet")
            _note(para, f'  [{ORIGIN_LABEL.get(item.get("origin"), item.get("origin", ""))}]')
            _translation(doc, item["text"], translations)
            if item.get("source_quote"):
                _note(doc.add_paragraph(), f'Source: "{item["source_quote"]}"')

    number += 1
    doc.add_heading(f"{number}. Functional requirements and traceability", 1)
    rules = brd.get("business_rules", [])
    if rules:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        for cell, title in zip(table.rows[0].cells, ["ID", "Requirement", "Spoken by / source", "Verification", "Approved"]):
            cell.text = title
        for r in rules:
            cells = table.add_row().cells
            cells[0].text = r.get("req_id", r["id"])
            translated = (translations or {}).get(r["rule"])
            cells[1].text = f'{r["title"]}\n{r["rule"]}' + (f"\n{translated}" if translated else "") + (
                "".join("\n• " + a for a in r.get("acceptance", [])))
            cells[2].text = (f'{r["speaker"]}: ' if r.get("speaker") else "") + f'"{r.get("source_quote", "")}"' + (
                f'\nAudio {r["start"]}s–{r["end"]}s' if r.get("start") is not None else "")
            check = r.get("verification") or {}
            cells[3].text = check.get("status", "").replace("_", " ") + "".join(
                f'\n{n["value"]}: {n["status"]}' + (f' ("{n["spoken"]}")' if n.get("spoken") and n["status"] == "interpreted" else "") for n in check.get("numbers", []))
            cells[4].text = f'{r.get("approved_by") or "Pending"}\n{r.get("approved_at") or ""}'
    else:
        _note(doc.add_paragraph(), "No approved requirements.")

    number += 1
    doc.add_heading(f"{number}. Open questions", 1)
    for q in brd.get("open_questions") or ["None."]:
        doc.add_paragraph(q, style="List Bullet")
        _translation(doc, q, translations)

    if brd.get("rejected"):
        number += 1
        doc.add_heading(f"{number}. Verification report: rejected AI content", 1)
        _note(doc.add_paragraph(), "These statements were drafted by the AI but blocked because they could not be traced to what was said.")
        for r in brd["rejected"]:
            para = doc.add_paragraph(style="List Bullet")
            run = para.add_run(r.get("text", ""))
            run.font.strike = True
            _note(para, "  " + r.get("reason", ""))

    number += 1
    doc.add_heading(f"{number}. Sign-off", 1)
    doc.add_paragraph(brd.get("authorization", ""))
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
