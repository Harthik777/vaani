"""BRD -> Excel requirements traceability matrix (RTM): one row per requirement, source and status in columns."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STATUS = {"verified": "Verified against speech", "needs_confirmation": "Numbers to confirm", "unsupported_number": "Unspoken number", "rejected": "Quote mismatch"}


def build_xlsx(brd: dict, translations: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Traceability"
    tr = translations or {}
    head = ["ID", "Title", "Requirement (IF / THEN / UNLESS)", "Translation", "Said by", "Exact words", "Clip (s)", "Verifier status", "Acceptance criteria", "Approved by", "Approved at", "Jira summary"]
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F6E56")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for r in brd.get("business_rules") or []:
        v = (r.get("verification") or {}).get("status", "")
        clip = f"{r['start']}–{r['end']}" if r.get("start") is not None else ""
        ws.append([r.get("req_id", ""), r.get("title", ""), r.get("rule", ""), tr.get(r.get("rule"), ""), r.get("speaker") or "", r.get("source_quote") or "", clip,
                   STATUS.get(v, v), "\n".join(r.get("acceptance") or []), r.get("approved_by") or "", r.get("approved_at") or "", f"{r.get('req_id', '')}: {r.get('title', '')}"])
    widths = [8, 28, 48, 36, 12, 48, 10, 20, 44, 22, 22, 36]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"

    oq = wb.create_sheet("Open questions")
    oq.append(["#", "Question"])
    for i, q in enumerate(brd.get("open_questions") or [], 1):
        oq.append([i, q])
    oq.column_dimensions["B"].width = 90

    sec = wb.create_sheet("Sections")
    sec.append(["Section", "Item", "Origin", "Exact words"])
    titles = brd.get("section_titles") or {}
    for key, items in (brd.get("sections") or {}).items():
        for item in items:
            sec.append([titles.get(key, key), item.get("text", ""), item.get("origin", ""), item.get("source_quote") or ""])
    for col, w in zip("ABCD", [26, 60, 14, 60]):
        sec.column_dimensions[col].width = w

    if brd.get("rejected"):
        rej = wb.create_sheet("Blocked by verifier")
        rej.append(["Text the AI wrote", "Reason"])
        for r in brd["rejected"]:
            rej.append([r.get("text", ""), r.get("reason", "")])
        rej.column_dimensions["A"].width = 70
        rej.column_dimensions["B"].width = 50

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
